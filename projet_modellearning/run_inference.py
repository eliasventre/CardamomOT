"""
CardamomOT on the runs of the benchmark networks (bench/, generate.py), with observed or hidden proteins, with and
without model learning of the GRN sigma, then evaluation: family of each target (model learning), network (AUPR,
sign of the true edges) and error on theta.

Usage: python run_inference.py <network> [runs] [--modes obs_ml,obs_base,hid_ml,hid_base] [--compare-only]
       (e.g. python run_inference.py FN4 1-3; runs '1-3' or '1,4,7', default 1)
Projects: bench/<network>/runs/run<r>_<mode>/; summary appended to bench/<network>/summary.csv.
"""
import argparse
import csv
import os
import shutil
import subprocess
import sys

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
BENCH = os.path.join(HERE, 'bench')
STEPS = ['split_dataset', 'select_genes', 'get_degradation_rates', 'infer_mixture', 'infer_network_structure']
# Inference options of the simulated model (proteins only, mRNA read at the observation: no bridge, no memory)
PARAMS = {'split': 'full', 'select_genes': False, 'build_prior_network': False, 'estimate_proliferation_rates': False,
          'simulate_with_proliferation': False, 'simulate_perturbations': False, 'run_test': False,
          'run_classical_OT': False, 'allow_depth_correction': False, 'compute_depth_factor': False,
          'inference_mrnas_direct': True, 'mrna_bridge': False, 'targets_accounting_late_mrna': False, 'seed': 0}
# obs_*: Data/data_prot.h5ad given (observed proteins); hid_*: proteins reconstructed by the trajectories
MODES = {'obs_ml': (True, {'model_learning_grn': True}), 'obs_base': (True, {'model_learning_grn': False}),
         'hid_ml': (False, {'model_learning_grn': True}), 'hid_base': (False, {'model_learning_grn': False})}


def setup(name, run, mode):
    proj = os.path.join(BENCH, name, 'runs', f'run{run}_{mode}')
    os.makedirs(os.path.join(proj, 'Data'), exist_ok=True)
    shutil.copy(os.path.join(BENCH, name, 'Data', f'data_{run}.h5ad'), os.path.join(proj, 'Data', 'data.h5ad'))
    observed, mode_params = MODES[mode]
    if observed:
        shutil.copy(os.path.join(BENCH, name, 'Data', f'data_prot_{run}.h5ad'), os.path.join(proj, 'Data', 'data_prot.h5ad'))
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(REPO, 'CardamomOT_inputs.xlsx'))
    ws = wb['model_parameters']
    hdr = [c.value for c in ws[1]]
    ip, iv = hdr.index('parameter'), hdr.index('value')
    rows = {row[ip].value: row for row in ws.iter_rows(min_row=2)}
    for k, v in {**PARAMS, **mode_params}.items():
        if k in rows:
            rows[k][iv].value = v
        else:
            ws.append([k, v] + [None] * (len(hdr) - 2))
    wb.save(os.path.join(proj, 'Data', 'CardamomOT_inputs.xlsx'))
    return proj


def run_steps(proj):
    for step in STEPS:
        print(f'=== {step}', flush=True)
        subprocess.run([sys.executable, '-m', 'CardamomOT.cli', 'step', step, '-i', proj], cwd=REPO, check=True)


def aupr(score, y):
    """Average precision of the scores against the binary labels y."""
    order = np.argsort(-score)
    tp = np.cumsum(y[order])
    prec = tp / np.arange(1, len(order) + 1)
    return float(np.sum(prec * y[order]) / max(y.sum(), 1))


def evaluate(name, run, mode, proj):
    from CardamomOT.inference.sigma import SPEC_PHI, SPEC_M, PHI_LOG, PHI_HERBACH, is_herbach
    truth = np.load(os.path.join(BENCH, name, 'True', f'truth_{run}.npz'))
    inter = np.load(os.path.join(proj, 'cardamomOT', 'inter.npy'))
    inter = inter[..., 0] if inter.ndim == 3 and inter.shape[-1] == 1 else inter
    T = np.sign(truth['inter'])
    G1 = T.shape[0]
    off = ~np.eye(G1, dtype=bool)
    off[:, 0] = False                        # no edge towards the stimulus
    y = (np.abs(T) > 0)[off]
    res = dict(network=name, run=run, mode=mode, aupr=aupr(np.abs(inter)[off], y), aupr_random=float(y.mean()))
    edges = np.argwhere((np.abs(T) > 0) & off)
    res['sign_ok'] = float(np.mean([np.sign(inter[i, j]) == T[i, j] for i, j in edges]))
    # Mean absolute error on theta (interactions towards the genes, basal of the genes)
    basal = np.atleast_2d(np.load(os.path.join(proj, 'cardamomOT', 'basal.npy')))[0]
    res['theta_mae'] = float(np.mean(np.abs(inter - truth['inter'])[:, 1:]))
    res['basal_mae'] = float(np.mean(np.abs(basal[1:] - truth['basal'][1:])))
    # True family of each gene (phi: 0 sigmoid, 1 Hill type, 2 Herbach type) and Hill coefficient of the Herbach ones
    fam = truth['family'].astype(int)
    fam = np.where([is_herbach(f) for f in fam], PHI_HERBACH, fam)  # both forms of the Herbach type compared alike
    m_true = truth['spec'][:, SPEC_M] if truth['spec'].shape[1] > SPEC_M else np.zeros(len(fam))
    spec_path = os.path.join(proj, 'cardamomOT', 'sigma_spec_grn.npy')
    if os.path.exists(spec_path):
        spec = np.load(spec_path)
        sel = spec[:, SPEC_PHI].astype(int)
        sel = np.where([is_herbach(f) for f in sel], PHI_HERBACH, sel)
        recall = lambda f: float(np.mean(sel[1:][fam[1:] == f] == f)) if (fam[1:] == f).any() else np.nan
        token = lambda f, m: {0: 'S', PHI_LOG: 'H'}.get(int(f), f'B{m:g}')
        both = (fam[1:] == PHI_HERBACH) & (sel[1:] == PHI_HERBACH)
        res.update(family_acc=float(np.mean(sel[1:] == fam[1:])), hill_recall=recall(PHI_LOG), sigmoid_recall=recall(0),
                   herbach_recall=recall(PHI_HERBACH),
                   herbach_m_acc=float(np.mean(np.isclose(spec[1:, SPEC_M][both], m_true[1:][both]))) if both.any() else np.nan,
                   true_family=' '.join(token(f, m) for f, m in zip(fam[1:], m_true[1:])),
                   selected=' '.join(token(f, m) for f, m in zip(sel[1:], spec[1:, SPEC_M])))
        res['exact_acc'] = float(np.mean([a == b for a, b in zip(res['true_family'].split(), res['selected'].split())]))
    np.set_printoptions(precision=2, suppress=True, linewidth=150)
    print(f'\n{name} run {run} [{mode}]: inferred inter (rows = regulators, 0 = stimulus)\n{inter}')
    print(', '.join(f'{k} {v:.2f}' if isinstance(v, float) else f'{k} {v}' for k, v in res.items()))
    return res


def append_summary(name, res):
    path = os.path.join(BENCH, name, 'summary.csv')
    keys = ['network', 'run', 'mode', 'aupr', 'aupr_random', 'sign_ok', 'theta_mae', 'basal_mae', 'family_acc',
            'hill_recall', 'sigmoid_recall', 'herbach_recall', 'herbach_m_acc', 'true_family', 'selected']
    new = not os.path.exists(path)
    with open(path, 'a', newline='') as f:
        w = csv.DictWriter(f, fieldnames=keys, extrasaction='ignore')
        if new:
            w.writeheader()
        w.writerow({k: res.get(k, '') for k in keys})


def parse_runs(s):
    if '-' in s:
        a, b = s.split('-')
        return list(range(int(a), int(b) + 1))
    return [int(x) for x in s.split(',')]


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('network')
    ap.add_argument('runs', nargs='?', default='1')
    ap.add_argument('--modes', default='obs_ml,obs_base,hid_ml,hid_base')
    ap.add_argument('--compare-only', action='store_true')
    a = ap.parse_args()
    sys.path.insert(0, REPO)
    for run in parse_runs(a.runs):
        for mode in a.modes.split(','):
            proj = setup(a.network, run, mode) if not a.compare_only else \
                os.path.join(BENCH, a.network, 'runs', f'run{run}_{mode}')
            if not a.compare_only:
                run_steps(proj)
            append_summary(a.network, evaluate(a.network, run, mode, proj))
