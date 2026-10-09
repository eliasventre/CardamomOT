"""
Optional run inputs of a project: one Excel workbook, Data/CardamomOT_inputs.xlsx, filled by the user
(empty template: CardamomOT_inputs.xlsx at the root of the repository; without workbook, every default
applies). Each sheet holds one kind of information (gene lists, stimulus schedules and targets,
perturbations, timepoints, proliferation and transition rates...); empty cells mean "not given".
The first call of input_dir(project) in a process exports the workbook to cardamomOT/inputs/ as the
files the pipeline reads (never written by hand). Large numeric arrays (reference_network.csv,
basal_init / basal_ref, inter_init / inter_ref, inter_simul_ref) stay files in Data/.
"""
import os
import re
import shutil

import numpy as np
import pandas as pd

WORKBOOK = 'CardamomOT_inputs.xlsx'
CACHE = os.path.join('cardamomOT', 'inputs')

# Parameters listed in the model_parameters sheet, grouped by use (empty value = default of base.py;
# other attributes of NetworkModel can be added as rows)
PARAMETER_GROUPS = [
    ('Pipeline: steps and data', [
        'estimate_proliferation_rates', 'simulate_with_proliferation', 'select_genes', 'build_prior_network',
        'simulate_perturbations', 'run_test', 'split', 'train_rate', 'species']),
    ('Gene selection', [
        'num_max_genes', 'n_query_genes', 'n_entropy_genes', 'network_method', 'sample_network_combination',
        'k_in_steiner', 'closure_min',
        'min_entropy_change', 'min_nb_separation']),
    ('Literature prior network', [
        'prior_network_pen', 'max_free_params', 'literature_selection', 'literature_depth', 'literature_resources']),
    ('Stimulus', ['stimulus']),
    ('Per-cell sequencing depth', [
        'use_depth_factor', 'compute_depth_factor', 'allow_depth_correction', 'depth_method', 'depth_by_cell_type']),
    ('NB mixture and samples', [
        'mean_forcing_em', 'batch_size_mixture', 'soft_em_refinement', 'integrate_samples', 'ref_sample_integration']),
    ('Network inference', [
        'force_basins', 'temporal_basins', 'min_n_loops', 'max_iter', 'batch_size_traj', 'n_network_fits',
        'constrain_basal_uniform', 'network_condition_pen']),
    ('Proliferation and simulation', [
        'growth_reg_source', 'prolif_uses_stimulus', 'simulation_stochastic', 'protein_dilution']),
    ('Report', ['classifier_method', 'embedding_method_visualization', 'cell_depth_for_representation']),
    ('Reproducibility', ['seed']),
]

# Sheet -> (columns with their description, purpose)
SHEETS = {
    'model_parameters': (
        [('parameter', 'Attribute of NetworkModel (CardamomOT/model/base.py); rows can be added for other attributes'),
         ('value', 'Value for this project (empty = default of base.py); the command-line options of the pipeline '
                   '(--stimulus, --prior, --mean-forcing, --force-basins, --temporal-basins) override it'),
         ('description', 'What the parameter does (comment of base.py)')],
        'Parameters of the model for this project, grouped by use: empty value = default of CardamomOT/model/base.py; '
        'a filled value overrides it, and the command-line options override the workbook.'),
    'gene_lists': (
        [('genes_queries', 'Genes of interest for the gene selection (select_genes with select_genes = True), one per row'),
         ('proliferation_signatures', 'Proliferation marker genes (get_proliferation_rates), replace the built-in list'),
         ('death_signatures', 'Death marker genes (get_proliferation_rates), replace the built-in list'),
         ('senescence_signatures', 'Senescence / arrest marker genes (get_proliferation_rates), replace the built-in list')],
        'Gene lists, one gene per row in each column. Any other column is a named gene list, usable as the '
        'targets (STIMk) or a RATE target of perturbation_inference and perturbation_simulation '
        '(exported as gene_list_<name>.txt).'),
    'stimulus_inference_schedule': (
        [('sample_id', 'Optional: dataset_id whose schedule these rows replace (empty or "all": default schedule '
                       'of every sample)'),
         ('time', 'Timepoint (optional: without times, rows follow the sorted timepoints of the data)'),
         ('stimulus_1', 'Value of stimulus 1 at this timepoint (default: 0 at the first timepoint, 1 after)')],
        'Schedule of the stimuli of the measured data (inference); one row per timepoint, one column per stimulus '
        '(add stimulus_2, stimulus_3... for several stimuli; without this sheet: one stimulus). Rows with a '
        'sample_id give the schedule of that sample only (e.g. an untreated control at 0).'),
    'stimulus_test_schedule': (
        [('sample_id', 'Optional: dataset_id whose schedule these rows replace (empty or "all": default)'),
         ('time', 'Timepoint (optional: without times, rows follow the sorted timepoints)'),
         ('stimulus_1', 'Value of inference stimulus 1 for the held-out cells (default: the inference schedule)')],
        'Schedule of the inference stimuli for the held-out cells (infer_test): test split and samples removed '
        'from the inference (perturbation_inference, remove_from_inference). Empty: the inference schedule.'),
    'stimulus_simulation_schedule': (
        [('scenario', 'Optional: name of an alternative schedule of the simulations (empty: default schedule); '
                      'perturbation_simulation chooses the scenarios of each condition (column schedules)'),
         ('sample_id', 'Optional: dataset_id whose schedule these rows replace (empty or "all": default; '
                       'default schedule only)'),
         ('time', 'Simulated timepoint; the value holds from this time on and applies to the simulated '
                  'intervals ending after it (optional: without times, rows follow the sorted simulated timepoints)'),
         ('stimulus_1', 'Inference stimulus 1 during the simulations (default: its inference schedule)'),
         ('STIM1', 'Perturbation stimulus STIM1 of perturbation_simulation (default: 0 at the first time, 1 after)'),
         ('comment', 'Free comment (ignored)')],
        'Schedules of the simulations: inference stimuli (stimulus_k, acting on the genes and, through the RATEk '
        'of perturbation_inference, on proliferation and death), then perturbation stimuli (STIMk, as in '
        'perturbation_simulation). Rows with a scenario name give alternative schedules (e.g. alternating '
        'inference and perturbation stimuli), simulated for the conditions that ask for them; a value applies '
        'to the simulated intervals ending at or after its time, so switching times should be simulated times.'),
    'perturbation_inference': (
        [('sample_id', 'dataset_id of a measured sample carrying genetic perturbations (KO / OV), or "all" '
                       'for the row describing the inference stimuli (STIMk / RATEk)'),
         ('KO', 'Genes knocked-out in this sample, comma-separated'),
         ('OV', 'Genes over-expressed in this sample, comma-separated'),
         ('remove_from_inference', 'True / 1: every cell of this sample is held out (data_test, not split), '
                                   'excluded from the gene selection and the inference; infer_test validates the '
                                   'model on it with its stimulus_test_schedule'),
         ('reference_sample', 'Removed sample without the first timepoint: dataset_id whose first-timepoint '
                              'training cells start its validation simulation'),
         ('STIM1', 'Row "all": possible direct targets of inference stimulus 1: a gene list of gene_lists or '
                   'comma-separated genes (empty: every gene)'),
         ('RATE1', 'Row "all": effect of inference stimulus 1 on the net proliferation rate (per hour), '
                   'TARGET:delta comma-separated, scaled by its schedule: TARGET = a cell type of '
                   'cell_type_proliferation (delta added to its rate of proliferation_rates, which is then the '
                   'rate without stimulus), or a gene list of gene_lists / GENE1+GENE2 / a gene (delta x mRNA '
                   'signature score in [0, 1])'),
         ('comment', 'Free comment (ignored)')],
        'Perturbations of the measured data, used by the inference: genetic perturbations of samples (KO / OV), '
        'samples held out for validation (remove_from_inference, reference_sample), '
        'and the inference stimuli (row sample_id = all: targets STIMk and effects on proliferation RATEk, '
        'k = column of stimulus_inference_schedule). Empty: one stimulus, every gene a possible target, no '
        'effect on the rates.'),
    'perturbation_simulation': (
        [('KO', 'Knocked-out genes, comma-separated; GENE-X for a partial KO of X%'),
         ('OV', 'Over-expressed genes, comma-separated; GENE-X for a partial OV of X%'),
         ('STIM1', "Perturbation stimulus 1: targets followed by + (activated) or - (inhibited), e.g. CHGA+STMN2-"),
         ('RATE1', "Effect of stimulus 1 (its schedule, STIM1 may be empty) on the net proliferation rate: "
                   "TARGET:delta, comma-separated; TARGET = a gene list of gene_lists, GENE1+GENE2..., a gene or "
                   "'all'; delta (per time unit) added to R of a cell at the maximal score, e.g. "
                   "ferroptosis_sensitive:-0.01 (needs the proliferation MLP)"),
         ('schedules', "Schedules simulated for this condition, comma-separated: 'default' and/or scenario "
                       "names of stimulus_simulation_schedule, 'all' = default + every scenario (empty: default)"),
         ('comment', 'Free comment (ignored)')],
        'In-silico perturbations simulated by simulate_network_KOV, one condition per row '
        '(add STIM2, STIM3... for several perturbation stimuli; their effects add up).'),
    'times': (
        [('times_inference', 'Inference restricted to the timepoints <= the largest value of this column'),
         ('times_simulation', 'Timepoints of the simulations (0 added if absent)')],
        'Timepoints, one per row.'),
    'proliferation_rates': (
        [('cell_type', 'Cell type (as in obs cell_type_proliferation, else cell_type_transition, else cell_type)'),
         ('net_rate_per_hour', 'Reference net proliferation rate (birth - death), in h^-1; without the inference '
                               'stimulus if perturbation_inference gives its effect on this cell type (RATEk)'),
         ('sample_id', 'Optional: dataset_id these anchors are for (empty or "all": default of the samples that '
                       'have none of their own)')],
        'Anchors of the net proliferation rates per cell type (fit_population_anchors, get_proliferation_rates); '
        'every cell type needs a value. Rows with a sample_id anchor that sample only (a sample with rows of its own '
        'ignores the default ones); a sample without anchors borrows those of the best-fitting sample.'),
    'population_sizes': (
        [('time', 'Timepoint'), ('population_size', 'Total population size at this timepoint (any unit: only the '
                                                    'ratios between times of a sample are used)'),
         ('sample_id', 'Optional: dataset_id of the sample (empty or "all": default of the samples without rows)')],
        'Population sizes per sample and time: they constrain the absolute growth in fit_population_anchors and in '
        'the growth estimated by optimal transport.'),
    'transition_rates': (
        [('from \\ to', 'Source cell type (rows) and target cell types (header): transition rates in h^-1 '
                        '(off-diagonal, 0 = no direct transition; diagonal ignored)')],
        'Cell-type transition rate matrix (Markov generator) constraining optimal transport through the '
        'probabilities expm(Q dt): first column = source cell types, header = target cell types '
        '(as in obs cell_type_transition, else cell_type). To give a matrix per sample, add a column sample_id '
        'right of the matrix (one value per row of the block, or one block per sample); rows without sample_id '
        'are the default of the samples without block of their own.'),
}

def _is_empty(v):
    return v is None or (isinstance(v, float) and np.isnan(v)) or (isinstance(v, str) and not v.strip())


def _clean(v):
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


# ---------------------------------------------------------------------------
# Workbook creation and access
# ---------------------------------------------------------------------------

def _style_header(ws, columns):
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill, Alignment
    for j, (name, desc) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=j, value=name)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F3864')
        c.alignment = Alignment(horizontal='center')
        if desc:
            c.comment = Comment(desc, 'CardamomOT', width=320, height=90)
        ws.column_dimensions[c.column_letter].width = max(16, len(name) + 4)
    ws.freeze_panes = 'A2'


def _fill_readme(ws):
    """Content of the readme sheet: purpose of the workbook and of each sheet."""
    from openpyxl.styles import Font, Alignment
    lines = [
        ('CardamomOT — optional inputs of the run', 'title'),
        ('Fill only what you need: empty cells mean "not given" and the defaults of CardamomOT apply. '
         'Each sheet holds one kind of information; the header cells carry a comment (hover them) '
         'explaining their content.', ''),
        ('Large numeric arrays stay files in Data/: reference_network.csv (structural prior), '
         'basal_init / basal_ref, inter_init / inter_ref (.npy or .csv), inter_simul_ref (.npy or .csv).', ''),
        ('', ''),
        ('Sheet', 'header'),
    ]
    for i, (text, kind) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        if kind == 'title':
            c.font = Font(bold=True, size=14)
        elif kind == 'header':
            for j, h in enumerate(['Sheet', 'Content'], start=1):
                ws.cell(row=i, column=j, value=h).font = Font(bold=True)
        else:
            c.alignment = Alignment(wrap_text=True, vertical='top')
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
            ws.row_dimensions[i].height = 45
    r = len(lines) + 1
    for name, (columns, purpose) in SHEETS.items():
        ws.cell(row=r, column=1, value=name).font = Font(bold=True)
        ws.cell(row=r, column=2, value=purpose).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 48
        r += 1
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 110


def create_workbook(path):
    """Empty workbook with the documented structure (readme sheet + one sheet per input): the template of the
    repository root (python -c "from CardamomOT.inputs import create_workbook; create_workbook('CardamomOT_inputs.xlsx')")."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'readme'
    _fill_readme(ws)
    for name, (columns, purpose) in SHEETS.items():
        _style_header(wb.create_sheet(name), columns)
    _fill_parameter_rows(wb['model_parameters'])
    wb.save(path)


def upgrade_workbook(path, out=None):
    """
    Brings an existing workbook to the current layout without touching the values of the kept sheets: the header
    sample_id of proliferation_rates and population_sizes, and the removal of the former sheet doubling_times. Saved in
    place (out=None) or as `out`. Returns what was changed.
    """
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = load_workbook(path)
    added = []
    for sheet, pos in (('proliferation_rates', 3), ('population_sizes', 3)):
        ws = wb[sheet] if sheet in wb.sheetnames else None
        if ws is not None and 'sample_id' not in [c.value for c in ws[1]]:
            doc = dict(SHEETS[sheet][0])['sample_id']
            c = ws.cell(row=1, column=max(pos, ws.max_column + 1), value='sample_id')
            c.font, c.fill, c.alignment = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F3864'), Alignment(horizontal='center')
            c.comment = Comment(doc, 'CardamomOT', width=320, height=90)
            ws.column_dimensions[c.column_letter].width = 16
            added.append(f'column sample_id of {sheet}')
    if 'doubling_times' in wb.sheetnames:
        del wb['doubling_times']
        added.append('sheet doubling_times removed (population constraints: proliferation_rates, transition_rates, '
                     'population_sizes)')
    if added:
        wb.save(out or path)
    return added


def _truthy(v):
    return not _is_empty(v) and str(_clean(v)).strip().lower() in ('1', 'true', 'yes', 'oui', 'x', 'vrai')


def parameter_docs():
    """{attribute: (default, comment)} of NetworkModel, comments read from base.py."""
    from .model.base import NetworkModel
    import inspect
    model = NetworkModel(1)
    src = inspect.getsource(NetworkModel.__init__)
    docs = {}
    for name, comment in re.findall(r'self\.(\w+)\s*=\s*[^#\n]*#\s*(.*)', src):
        docs.setdefault(name, comment.strip())
    return {name: (getattr(model, name, None), docs.get(name, '')) for name in vars(model)}


def _fill_parameter_rows(ws, values=None, extra=()):
    """Group title rows, then parameter / value / description of each listed parameter (+ extra rows)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    docs, values = parameter_docs(), values or {}
    groups = PARAMETER_GROUPS + ([('Other parameters', list(extra))] if extra else [])
    r = 2
    for title, names in groups:
        for j in (1, 2, 3):
            ws.cell(row=r, column=j).fill = PatternFill('solid', fgColor='D9E1F2')
        ws.cell(row=r, column=1, value=title).font = Font(bold=True)
        r += 1
        for name in names:
            ws.cell(row=r, column=1, value=name)
            if name in values:
                ws.cell(row=r, column=2, value=values[name])
            ws.cell(row=r, column=3, value=docs.get(name, (None, ''))[1]).alignment = Alignment(wrap_text=True,
                                                                                               vertical='top')
            r += 1
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 120


def read_sheet(wb, name):
    """DataFrame of a sheet (header row 1), empty cells as NaN, empty rows dropped."""
    if name not in wb.sheetnames:
        return pd.DataFrame()
    rows = list(wb[name].iter_rows(values_only=True))
    if not rows:
        return pd.DataFrame()
    header = [str(h).strip() if h is not None else f'_col{j}' for j, h in enumerate(rows[0])]
    df = pd.DataFrame([list(r) + [None] * (len(header) - len(r)) for r in rows[1:]], columns=header)
    df = df.loc[:, [not h.startswith('_col') or df[h].notna().any() for h in header]]
    return df.dropna(how='all').reset_index(drop=True)


def write_sheet(wb, name, df):
    """Replace the content of a sheet by df (header kept and styled for the known columns)."""
    ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
    ws.delete_rows(1, ws.max_row)
    known = dict(SHEETS.get(name, ([], ''))[0])
    _style_header(ws, [(str(c), known.get(str(c), '')) for c in df.columns])
    for i, row in enumerate(df.itertuples(index=False), start=2):
        for j, v in enumerate(row, start=1):
            if not _is_empty(v):
                ws.cell(row=i, column=j, value=_clean(v))


# ---------------------------------------------------------------------------
# Workbook -> text files read by the pipeline
# ---------------------------------------------------------------------------

def _values(df, col):
    return [_clean(v) for v in df[col] if not _is_empty(v)] if col in df else []


def _schedule(df, cols):
    """Matrix of the schedule columns, rows sorted by time when every row has one."""
    df = df.dropna(subset=cols, how='all')
    if 'time' in df and df['time'].notna().all() and len(df):
        df = df.sort_values('time')
    return df[cols].astype(float).to_numpy()


def _split_samples(df):
    """(default rows, {sample_id: rows}) of a schedule sheet (sample_id empty or "all" = default)."""
    if 'sample_id' not in df or not len(df):
        return df, {}
    sid = df['sample_id'].astype(str).str.strip()
    default = df['sample_id'].isna() | sid.str.lower().isin(['', 'all', 'nan', 'none'])
    return df[default], {s: df[~default & (sid == s)] for s in sorted(set(sid[~default]))}


def _add_overrides(overrides, kind, groups, cols):
    """Per-sample schedules of a sheet into overrides[kind] (missing values: 1)."""
    for sid, rows in groups.items():
        rows = rows.dropna(subset=cols, how='all') if cols else rows.iloc[:0]
        if not len(rows):
            continue
        timed = 'time' in rows and rows['time'].notna().all()
        if timed:
            rows = rows.sort_values('time')
        overrides.setdefault(kind, {})[sid] = {
            'time': [float(t) for t in rows['time']] if timed else None,
            'values': rows[cols].astype(float).fillna(1.0).to_numpy().tolist()}


def export(wb, out_dir):
    """Write the content of the workbook as the text files read by the pipeline."""
    if os.path.isdir(out_dir):
        shutil.rmtree(out_dir)
    os.makedirs(out_dir)
    write = lambda name, text: open(os.path.join(out_dir, name), 'w').write(text)

    gl = read_sheet(wb, 'gene_lists')
    for col, name in [('genes_queries', 'genes_queries.txt'), ('proliferation_signatures', 'proliferation_signatures.txt'),
                      ('death_signatures', 'death_signatures.txt'), ('senescence_signatures', 'senescence_signatures.txt')]:
        genes = _values(gl, col)
        if genes:
            write(name, '\n'.join(map(str, genes)) + '\n')
    known = {'genes_queries', 'proliferation_signatures', 'death_signatures', 'senescence_signatures'}
    for col in gl.columns:
        genes = _values(gl, col) if col not in known and not str(col).startswith('_col') else []
        if genes:
            write(f'gene_list_{col}.txt', '\n'.join(map(str, genes)) + '\n')

    overrides = {}
    st_all = read_sheet(wb, 'stimulus_inference_schedule')
    inf_cols = sorted([c for c in st_all.columns if c.startswith('stimulus_') and st_all[c].notna().any()],
                      key=lambda c: int(c.split('_')[1]))
    st, st_over = _split_samples(st_all)
    if inf_cols and len(st.dropna(subset=inf_cols, how='all')):
        np.savetxt(os.path.join(out_dir, 'stimulus_schedule_inference.txt'), _schedule(st, inf_cols), fmt='%g')
    n_inf = max(1, len(inf_cols))
    _add_overrides(overrides, 'inference', st_over, inf_cols)

    te_all = read_sheet(wb, 'stimulus_test_schedule')
    te_cols = sorted([c for c in te_all.columns if c.startswith('stimulus_') and te_all[c].notna().any()],
                     key=lambda c: int(c.split('_')[1]))
    te, te_over = _split_samples(te_all)
    if te_cols and len(te.dropna(subset=te_cols, how='all')):
        np.savetxt(os.path.join(out_dir, 'stimulus_schedule_test.txt'), _schedule(te, te_cols), fmt='%g')
    _add_overrides(overrides, 'test', te_over, te_cols)

    # Inference stimuli (row all of perturbation_inference): targets STIMk and rate effects RATEk
    pi = read_sheet(wb, 'perturbation_inference')
    is_all = (pi['sample_id'].astype(str).str.strip().str.lower() == 'all') if 'sample_id' in pi else pd.Series([], dtype=bool)
    lists, rates = {}, {}
    for _, r in pi[is_all].iterrows():
        for c in pi.columns:
            m = re.fullmatch(r'(STIM|RATE)(\d+)', str(c))
            if not m or _is_empty(r[c]):
                continue
            k, cell = int(m.group(2)), str(_clean(r[c]))
            if m.group(1) == 'STIM':
                # A gene list of gene_lists, else comma-separated genes
                genes = _values(gl, cell) if cell in gl.columns else [g.strip() for g in re.split(r'[,\s]+', cell) if g.strip()]
                lists.setdefault(k, []).extend(str(g) for g in genes)
            else:
                from .tools.perturbations import parse_rate
                rates.setdefault(k, []).extend(parse_rate(cell))
    bad = sorted(k for k in set(lists) | set(rates) if k > n_inf)
    if bad:
        print(f"[CardamomOT] Warning: perturbation_inference describes inference stimuli {bad} but "
              f"stimulus_inference_schedule has {n_inf} stimulus column(s): ignored")
    if any(lists.get(k) for k in range(1, n_inf + 1)):
        cols = [lists.get(k, []) for k in range(1, n_inf + 1)]
        n = max(len(l) for l in cols)
        write('stimulus_targets.txt', '\n'.join('\t'.join(l[i] if i < len(l) else '' for l in cols)
                                                for i in range(n)) + '\n')
    rates = {k: v for k, v in rates.items() if k <= n_inf and v}
    if rates:
        import json
        json.dump({str(k): v for k, v in rates.items()}, open(os.path.join(out_dir, 'stimulus_rates.json'), 'w'),
                  indent=1)

    ss_all = read_sheet(wb, 'stimulus_simulation_schedule')
    s_inf = sorted([c for c in ss_all.columns if c.startswith('stimulus_') and ss_all[c].notna().any()],
                   key=lambda c: int(c.split('_')[1]))
    s_pert = sorted([c for c in ss_all.columns if re.fullmatch(r'STIM\d+', c) and ss_all[c].notna().any()],
                    key=lambda c: int(c[4:]))
    if 'scenario' in ss_all:
        scen = ss_all['scenario'].map(lambda v: '' if _is_empty(v) else str(_clean(v)))
    else:
        scen = pd.Series([''] * len(ss_all), index=ss_all.index, dtype=str)
    ss, ss_over = _split_samples(ss_all[scen == ''])
    _add_overrides(overrides, 'simulation', ss_over, s_inf + s_pert)
    bad = [c for c in ss_all.columns if re.fullmatch(r'stimulus_\d+', c) and int(c.split('_')[1]) > n_inf
           and ss_all[c].notna().any()]
    if bad:
        print(f"[CardamomOT] Warning: stimulus_simulation_schedule columns {bad} beyond the {n_inf} inference "
              f"stimuli of stimulus_inference_schedule: ignored")
        s_inf = [c for c in s_inf if c not in bad]

    def sim_matrix(rows, what):
        # (values: inference then perturbation stimuli, row times or None) of a simulation schedule
        rows = rows.dropna(subset=s_inf + s_pert, how='all')
        timed = 'time' in rows and rows['time'].notna().all() and len(rows) > 0
        if timed:
            rows = rows.sort_values('time')
        if not s_inf:
            # Inference stimuli not given: default schedule (0 at the first time, 1 after)
            print(f"[CardamomOT] Warning: {what} has perturbation stimuli but no inference stimulus column: "
                  f"default 0 at the first time and 1 after for the {n_inf} inference stimuli")
            inf = np.ones((len(rows), n_inf))
            inf[0] = 0
        else:
            inf = rows[s_inf].astype(float).ffill().fillna(1.0).to_numpy()
        pert = rows[s_pert].astype(float).fillna(1.0).to_numpy() if s_pert else np.zeros((len(rows), 0))
        return np.hstack([inf, pert]), ([float(t) for t in rows['time']] if timed else None)

    if (s_inf or s_pert) and len(ss.dropna(subset=s_inf + s_pert, how='all')):
        vals, times = sim_matrix(ss, 'sheet stimulus_simulation_schedule')
        np.savetxt(os.path.join(out_dir, 'stimulus_schedule_simulate.txt'), vals, fmt='%g')
        if times is not None:
            write('stimulus_schedule_simulate_times.txt', '\n'.join(f'{t:g}' for t in times) + '\n')
    # Alternative scenarios: {name: {"time": [...] or null, "values": [[...]], "n_inference": n}}
    scenarios = {}
    for name in sorted(set(scen) - {''}):
        rows = ss_all[scen == name]
        default_rows, per_sample = _split_samples(rows)
        if per_sample:
            print(f"[CardamomOT] Warning: scenario '{name}' of stimulus_simulation_schedule: per-sample rows "
                  f"{sorted(per_sample)} ignored (scenarios have one schedule for every sample)")
        if not (s_inf or s_pert) or not len(default_rows.dropna(subset=s_inf + s_pert, how='all')):
            print(f"[CardamomOT] Warning: scenario '{name}' of stimulus_simulation_schedule has no value: ignored")
            continue
        vals, times = sim_matrix(default_rows, f"scenario '{name}' of stimulus_simulation_schedule")
        scenarios[name] = {'time': times, 'values': vals.tolist(), 'n_inference': n_inf}
    if scenarios:
        import json
        json.dump(scenarios, open(os.path.join(out_dir, 'simulation_scenarios.json'), 'w'), indent=1)

    pt = read_sheet(wb, 'perturbation_simulation')
    pcols = [c for c in pt.columns if c in ('KO', 'OV', 'schedules') or re.fullmatch(r'(STIM|RATE)\d+', str(c))]
    if pcols and len(pt):
        lines = ['\t'.join(pcols)]
        for _, r in pt.iterrows():
            cells = ['0' if _is_empty(r[c]) else str(_clean(r[c])) for c in pcols]
            if any(v != '0' for c, v in zip(pcols, cells) if c != 'schedules'):
                lines.append('\t'.join(cells))
        if len(lines) > 1:
            write('KO_OV_Stim_simulate.txt', '\n'.join(lines) + '\n')

    ki = pi[~is_all] if len(pi) else pi
    if 'sample_id' in ki and len(ki):
        lines = ['SAMPLE_ID\tKO\tOV']
        for _, r in ki.iterrows():
            if not _is_empty(r['sample_id']) and (not _is_empty(r.get('KO')) or not _is_empty(r.get('OV'))):
                lines.append('\t'.join(['0' if _is_empty(r.get(c)) else str(_clean(r.get(c)))
                                        for c in ('sample_id', 'KO', 'OV')]))
        if len(lines) > 1:
            write('KO_OV_inference.txt', '\n'.join(lines) + '\n')

    # Samples held out of the inference (validation in infer_test), with their reference sample
    if 'sample_id' in ki and 'remove_from_inference' in ki and len(ki):
        info = {'removed': [], 'reference': {}}
        for _, r in ki.iterrows():
            if not _is_empty(r['sample_id']) and _truthy(r['remove_from_inference']):
                sid = str(_clean(r['sample_id']))
                info['removed'].append(sid)
                if not _is_empty(r.get('reference_sample')):
                    info['reference'][sid] = str(_clean(r['reference_sample']))
        if info['removed']:
            import json
            json.dump(info, open(os.path.join(out_dir, 'samples_info.json'), 'w'), indent=1)

    tm = read_sheet(wb, 'times')
    old = [c for c in ('times_to_inference', 'times_to_simulate') if c in tm.columns and tm[c].notna().any()]
    if old:
        print(f"[CardamomOT] Warning: sheet times: columns {old} ignored (now times_inference, times_simulation)")
    for col in ('times_inference', 'times_simulation'):
        vals = _values(tm, col)
        if vals:
            write(f'{col}.txt', '\n'.join(f'{float(v):g}' for v in vals) + '\n')

    # Population sizes: rows without sample_id = default, the others per sample
    ps = read_sheet(wb, 'population_sizes')
    if all(c in ps for c in ('time', 'population_size')):
        ps = ps.dropna(subset=['time', 'population_size'])
        if len(ps):
            import json
            default, per_sample = _split_samples(ps)
            entry = lambda d: {f'{float(_clean(t)):g}': float(_clean(n)) for t, n in zip(d['time'], d['population_size'])}
            spec = {'default': entry(default) if len(default) else None,
                    'samples': {sid: entry(d) for sid, d in per_sample.items() if len(d)}}
            json.dump(spec, open(os.path.join(out_dir, 'population_sizes.json'), 'w'), indent=1)

    # Proliferation anchors: rows without sample_id = default (proliferation_rates.txt), the others per sample (json)
    pr = read_sheet(wb, 'proliferation_rates')
    if all(c in pr for c in ('cell_type', 'net_rate_per_hour')):
        pr = pr.dropna(subset=['cell_type', 'net_rate_per_hour'])
        default, per_sample = _split_samples(pr)
        if len(default):
            write('proliferation_rates.txt', '\n'.join(f'{_clean(a)}\t{_clean(b)}' for a, b in
                                                      zip(default['cell_type'], default['net_rate_per_hour'])) + '\n')
        if per_sample:
            import json
            json.dump({sid: {str(_clean(a)): float(b) for a, b in zip(d['cell_type'], d['net_rate_per_hour'])}
                       for sid, d in per_sample.items()}, open(os.path.join(out_dir, 'proliferation_rates_samples.json'), 'w'), indent=1)

    if overrides:
        import json
        json.dump(overrides, open(os.path.join(out_dir, 'stimulus_schedules.json'), 'w'), indent=1)

    mp = read_sheet(wb, 'model_parameters')
    if 'parameter' in mp and 'value' in mp:
        import json
        values = {str(k).strip(): _clean(v) for k, v in zip(mp['parameter'], mp['value'])
                  if not _is_empty(k) and not _is_empty(v)}
        if values:
            json.dump(values, open(os.path.join(out_dir, 'model_parameters.json'), 'w'), indent=1, default=str)

    tr = read_sheet(wb, 'transition_rates')
    if len(tr) and tr.shape[1] > 1:
        first = tr.columns[0]
        default, per_sample = _split_samples(tr)
        for sid, d in [(None, default)] + list(per_sample.items()):
            if not len(d):
                continue
            m = d.drop(columns=[c for c in ('sample_id',) if c in d]).set_index(first)
            m.index = m.index.astype(str)
            m = m.dropna(axis=1, how='all')
            m.to_csv(os.path.join(out_dir, 'transition_rates.csv' if sid is None else f'transition_rates__{sid}.csv'))


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

_SYNCED = {}


def removed_samples(project, present=None):
    """(removed dataset_id, {removed: reference dataset_id}) of perturbation_inference; with `present` (dataset_id of
    the data), the samples absent from the data are ignored with a warning."""
    import json
    path = os.path.join(input_dir(project), 'samples_info.json')
    if not os.path.exists(path):
        return [], {}
    info = json.load(open(path))
    removed = [str(s) for s in info.get('removed', [])]
    ref = {str(k): str(v) for k, v in info.get('reference', {}).items()}
    if present is not None:
        present = {str(s) for s in present}
        absent = [s for s in removed if s not in present]
        if absent:
            print(f"[CardamomOT] Warning: remove_from_inference given for sample(s) {absent} absent from the data: ignored")
        removed = [s for s in removed if s in present]
        for s in [s for s in removed if s in ref and ref[s] not in present]:
            print(f"[CardamomOT] Warning: reference_sample '{ref[s]}' of sample {s} absent from the data: ignored")
            del ref[s]
        ref = {s: r for s, r in ref.items() if s in removed}
    return removed, ref


def proliferation_sample_anchors(project):
    """{dataset_id: {cell type: net rate (h^-1)}} of the rows of proliferation_rates with a sample_id ({} if none)."""
    import json
    path = os.path.join(input_dir(project), 'proliferation_rates_samples.json')
    return json.load(open(path)) if os.path.exists(path) else {}


def population_sizes(project):
    """({time: size} or None default, {dataset_id: {time: size}}) of the population_sizes sheet."""
    import json
    path = os.path.join(input_dir(project), 'population_sizes.json')
    if not os.path.exists(path):
        return None, {}
    spec = json.load(open(path))
    conv = lambda d: {float(t): float(n) for t, n in d.items()} if d else None
    return conv(spec.get('default')), {s: conv(d) for s, d in spec.get('samples', {}).items()}


def sample_population_sizes(project, sample):
    """{time: size} of a sample (its rows, else the default rows), None if none."""
    default, per = population_sizes(project)
    return per.get(str(sample)) or default


def fitted_anchors(project):
    """Corrected anchors of fit_population_anchors.py (cardamomOT/population_anchors.json), None if absent."""
    import json
    path = os.path.join(os.path.abspath(project), 'cardamomOT', 'population_anchors.json')
    return json.load(open(path)) if os.path.exists(path) else None


def load_transition_rates(project, fitted=True):
    """
    Transition rate matrices: None, the default DataFrame (no per-sample block), or {'default': DataFrame or None,
    dataset_id: DataFrame} (for NetworkModel._load_ot_constraints). Those of fit_population_anchors.py (corrected on
    the proportions, per sample) if it was run and fitted some, else those of the transition_rates sheet (fitted=False:
    always the sheet).
    """
    import glob
    anchors = fitted_anchors(project) if fitted else None
    if anchors and anchors.get('transitions'):
        out = {'default': None}
        for sid, t in anchors['transitions'].items():
            m = pd.DataFrame(0.0, index=t['types'], columns=t['types'])
            for a_, row in t['matrix'].items():
                for b_, v in row.items():
                    m.loc[a_, b_] = v
            out[sid] = m
        return out
    d = input_dir(project)
    read = lambda f: pd.read_csv(f, sep=None, engine='python', index_col=0).rename(index=str).rename(columns=str)
    default = read(os.path.join(d, 'transition_rates.csv')) if os.path.exists(os.path.join(d, 'transition_rates.csv')) else None
    per = {os.path.basename(f)[len('transition_rates__'):-4]: read(f) for f in sorted(glob.glob(os.path.join(d, 'transition_rates__*.csv')))}
    return {'default': default, **per} if per else default


def project_parameters(project):
    """{attribute: value} filled in the model_parameters sheet of the project."""
    import json
    path = os.path.join(input_dir(project), 'model_parameters.json')
    return json.load(open(path)) if os.path.exists(path) else {}


def depth_factor_used(project):
    """use_depth_factor of the project (model_parameters sheet, else the default of NetworkModel)."""
    from .model.base import NetworkModel
    m = NetworkModel(1)
    m.apply_project_parameters(project, verb=False)
    return bool(m.use_depth_factor)


def input_dir(project):
    """
    Directory of the run inputs of a project (cardamomOT/inputs/), exported once per process from
    Data/CardamomOT_inputs.xlsx (empty without workbook: every default applies).
    """
    project = os.path.abspath(project)
    if project in _SYNCED:
        return _SYNCED[project]
    from openpyxl import load_workbook
    path = os.path.join(project, 'Data', WORKBOOK)
    out = os.path.join(project, CACHE)
    if os.path.exists(path):
        export(load_workbook(path), out)
    else:
        if os.path.isdir(out):
            shutil.rmtree(out)
        os.makedirs(out)
        print(f"[CardamomOT] No Data/{WORKBOOK}: defaults for every optional input (empty template: "
              f"{WORKBOOK} at the root of the repository)")
    _SYNCED[project] = out
    return out
